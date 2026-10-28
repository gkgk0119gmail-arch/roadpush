"""성과보고서 PPT(공식 양식 기반 5p) 자동 생성. 수치는 평가 JSON 에서 읽는다.

    python scripts/build_report.py --out dist/RoadPush_성과보고서.pptx [--pdf]

공고 요구(탐구참여인력·문제정의·AI 접근방법·결과 해석·한계 및 개선방안)를 양식 7개 항목에 배치:
  p1 1. AI 기술명 + 탐구참여인력      p2 2. 아이디어명 / 문제 정의 / 3. 기술요약
  p3 4. 최종목표 / 5-① AI 접근방법     p4 5-② 결과 해석
  p5 6. 애로점 / 한계 및 개선방안 / 7. 기대효과
양식(docs/templates/성과보고서_양식_PPT_5p.pptx = 공식 4p 양식 + 4쪽 복제)의 머리글·로고는 그대로 쓰고,
도식은 PPT 도형으로 직접 그려 편집 가능하게 한다. 그림은 scripts/make_figures.py 결과.
"""
from __future__ import annotations

import argparse
import json
import subprocess
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Cm, Pt

ROOT = Path(__file__).resolve().parent.parent
FIG, MEDIA = ROOT / "docs" / "figures", ROOT / "docs" / "media"
MLY = ROOT / "outputs" / "mly_downtown"
FONT, FONT_B = "Pretendard", "Pretendard SemiBold"


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


INK, INK2, MUTED = rgb("16181D"), rgb("4D5059"), rgb("8A8D96")
EMPH, EMPH_BG, CARD = rgb("2A78D6"), rgb("E8F1FC"), rgb("F4F6F9")
LINE, WHITE = rgb("D5D9E0"), rgb("FFFFFF")
RED, AMBER, GREEN = rgb("D03B3B"), rgb("F0A30A"), rgb("1E9E4A")
L, R = 3.5, 64.2                      # 본문 좌우 경계(cm)
CW = R - L


def _j(rel: str) -> dict:
    p = ROOT / rel
    return json.loads(p.read_text()) if p.exists() else {}


def numbers() -> dict:
    surf = _j("outputs/eval/eval_surface_final_full_v2.json")
    s150 = _j("outputs/eval/eval_surface_final_150_v2.json")
    sam3 = _j("outputs/teachers/eval_sam3_surface_v2.json")
    obs = _j("outputs/eval/eval_obstacle_v2_full.json") or _j("outputs/eval/eval_obstacle_final_full.json")
    push = _j("docs/figures/push_stats.json")
    f = lambda d, k, fmt: fmt.format(d[k]) if k in d else "평가 중"  # noqa: E731
    n = {"ratio_mae": f(surf, "ratio_mae", "{:.3f}"), "within": f(surf, "ratio_within_0.1", "{:.0%}"),
         "vanish_f1": f(surf, "vanish_f1", "{:.3f}"), "surf_miou": f(surf, "miou", "{:.2f}"),
         "sidewalk_iou": f"{surf['iou']['sidewalk']:.2f}" if surf else "평가 중",
         "n_surf": f"{surf.get('n_images', 0):,}", "n_obs": f"{obs.get('n_images', 0):,}",
         "mae150": f(s150, "ratio_mae", "{:.3f}"), "sam3_mae150": f(sam3, "ratio_mae", "{:.3f}"),
         "obs_miou": f(obs, "miou", "{:.3f}")}
    n["x_better"] = f"{sam3['ratio_mae'] / s150['ratio_mae']:.1f}배" if s150 and sam3 else "-"
    n.update(links=push["measured"], km=push["measured_km"], push_links=push["links"], push_km=push["km"],
             shared=push["shared"], frames=f"{push['frames']:,}", grades=push["grades"])
    return n


# ---------------------------------------------------------------- 그리기 도우미
def text(slide, x, y, w, h, paras, size=19, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space=0.35, margin=0.0):
    """paras: 문단 목록. 문단 = 문자열 또는 [(문자열, {size,color,bold}) ...] 런 목록."""
    tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.word_wrap, tf.vertical_anchor = True, anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Cm(margin)
    for i, p in enumerate(paras if isinstance(paras, list) else [paras]):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        runs = [(p, {})] if isinstance(p, str) else p
        sz = max(o.get("size", size) for _, o in runs)
        para.space_after = Pt(sz * space)
        for t, o in runs:
            r = para.add_run()
            r.text = t
            b = o.get("bold", bold)
            r.font.size, r.font.bold = Pt(o.get("size", size)), b
            r.font.color.rgb = o.get("color", color)
            r.font.name = FONT_B if b else FONT
    return tb


def box(slide, x, y, w, h, fill=CARD, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06, dash=False):
    sh = slide.shapes.add_shape(shape, Cm(x), Cm(y), Cm(w), Cm(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb, sh.line.width = line, Pt(1.5)
        if dash:
            from pptx.enum.dml import MSO_LINE
            sh.line.dash_style = MSO_LINE.DASH
    sh.shadow.inherit = False
    return sh


def badge(slide, x, y, label, d=1.25, fill=EMPH, size=17):
    c = box(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label
    r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), True, WHITE, FONT_B
    return c


def header(slide, x, y, title, w=30, sub=None):
    """섹션 머리: 파란 점 + 굵은 제목(+ 회색 부제)."""
    box(slide, x, y + 0.42, 0.55, 0.55, fill=EMPH, shape=MSO_SHAPE.OVAL)
    runs = [(title, {"size": 24, "bold": True})]
    if sub:
        runs.append(("   " + sub, {"size": 17, "color": INK2}))
    text(slide, x + 0.95, y, w - 0.95, 1.4, [runs], anchor=MSO_ANCHOR.MIDDLE)


def arrow(slide, x1, y1, x2, y2, color=INK2, dash=False, width=2.0):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    ln.line.color.rgb, ln.line.width = color, Pt(width)
    if dash:
        from pptx.enum.dml import MSO_LINE
        ln.line.dash_style = MSO_LINE.DASH
    ln_el = ln.line._get_or_add_ln()
    from pptx.oxml.ns import qn
    tail = ln_el.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln_el.append(tail)
    return ln


def picture(slide, path: Path, x, y, w, h=None, cover=True):
    """w×h 상자를 꽉 채우도록(가운데 기준 잘라내기) 넣는다. h 없으면 원본 비율."""
    iw, ih = Image.open(path).size
    if h is None:
        return slide.shapes.add_picture(str(path), Cm(x), Cm(y), Cm(w), Cm(w * ih / iw))
    pic = slide.shapes.add_picture(str(path), Cm(x), Cm(y), Cm(w), Cm(h))
    if cover:
        box_r, img_r = w / h, iw / ih
        if img_r > box_r:
            c = (1 - box_r / img_r) / 2
            pic.crop_left = pic.crop_right = c
        else:
            c = (1 - img_r / box_r) / 2
            pic.crop_top = pic.crop_bottom = c
    pic.line.color.rgb, pic.line.width = LINE, Pt(0.75)
    return pic


def caption(slide, x, y, w, t):
    text(slide, x, y, w, 1.0, [t], size=15, color=INK2)


def set_title(slide, t: str, page: int | None = None):
    for sh in list(slide.shapes):
        if not sh.has_text_frame:
            continue
        s = sh.text_frame.text.strip()
        if s.startswith("※"):
            sh._element.getparent().remove(sh._element)
        elif sh.name == "TextBox 2":
            sh.width = Cm(55)
            p = sh.text_frame.paragraphs[0]
            p.runs[0].text = t
            for r in p.runs[1:]:
                r._r.getparent().remove(r._r)
        elif sh.name == "TextBox 12" and page:
            sh.text_frame.paragraphs[0].runs[0].text = str(page)


def bar_chart(slide, x, y, w, h, title, cats, vals, fmt, best_high: bool):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series(title, vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Cm(x), Cm(y), Cm(w), Cm(h), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = True
    ch.chart_title.text_frame.text = title
    tr = ch.chart_title.text_frame.paragraphs[0].runs[0]
    tr.font.size, tr.font.bold, tr.font.name, tr.font.color.rgb = Pt(18), True, FONT_B, INK
    ch.font.size, ch.font.name, ch.font.color.rgb = Pt(15), FONT, INK2
    va = ch.value_axis
    va.visible = False
    va.has_major_gridlines = False
    va.maximum_scale = max(vals) * 1.25
    va.minimum_scale = 0
    ca = ch.category_axis
    ca.format.line.color.rgb = LINE
    ca.tick_labels.font.size = Pt(15)
    pl = ch.plots[0]
    pl.gap_width = 70
    pl.has_data_labels = True
    dl = pl.data_labels
    dl.number_format, dl.number_format_is_linked = fmt, False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size, dl.font.bold, dl.font.color.rgb = Pt(18), True, INK
    ser = pl.series[0]
    for i in range(len(vals)):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = EMPH if i == len(vals) - 1 else rgb("B9BDC6")
    return ch


# ---------------------------------------------------------------- 페이지
def page1(s, n):
    set_title(s, "AI 기술명")   # 양식 1쪽 제목은 자동 번호 "1."
    text(s, L, 5.6, 31, 3.2, ["RoadPush"], size=60, bold=True)
    text(s, L, 9.0, 31, 1.6, ["보행 영상 기반 보행공간 단절 위험도 자동 산출 시스템"], size=26, bold=True)
    text(s, L, 10.9, 30.5, 3.6, ["거리 사진을 AI로 분할해 '실제로 지나갈 수 있는 폭'을 재고, 도로 구간(링크)마다 "
                                 "걷는 사람이 차도로 밀려나는 위험을 점수로 보여 주는 지자체용 보도 점검 도구"],
         size=19, color=INK2, space=0.2)

    header(s, L, 14.9, "탐구참여인력")
    rows = [("성명", "소속", "역할"), ("남윤상", "연세대학교", "기획 · 데이터 구축 · 모델 학습 · 웹앱 개발 (단독)")]
    tbl = s.shapes.add_table(2, 3, Cm(L), Cm(16.6), Cm(30.5), Cm(3.0)).table
    for c, wcm in enumerate((6.0, 7.5, 17.0)):
        tbl.columns[c].width = Cm(wcm)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid(); cell.fill.fore_color.rgb = EMPH_BG if i == 0 else WHITE
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.text = v
            r = cell.text_frame.paragraphs[0].runs[0]
            r.font.size, r.font.name, r.font.color.rgb = Pt(17), FONT_B if i == 0 else FONT, INK

    stats = [(n["ratio_mae"], "유효보행폭 비율 오차(MAE)", f"AI Hub 검증 {n['n_surf']}장, ±0.1 이내 {n['within']}"),
             (n["x_better"], "범용 AI 대비 오차 감소", f"SAM 3 zero-shot {n['sam3_mae150']} → {n['mae150']}"),
             (f"{n['push_links']}개", "차도로 밀려나는 링크", f"서울 도심 {n['links']}개 링크 중 · {n['push_km']} km"),
             ("서버 0대", "브라우저 안에서 AI 추론", "설치 없이 웹 접속 · 사진 외부 전송 없음")]
    cw, ch_, gx, gy, y0 = 14.85, 6.6, 0.8, 0.7, 22.0
    for i, (v, t, sub) in enumerate(stats):
        x, y = L + (i % 2) * (cw + gx), y0 + (i // 2) * (ch_ + gy)
        box(s, x, y, cw, ch_, fill=CARD)
        text(s, x + 0.8, y + 0.6, cw - 1.6, 2.6, [v], size=40, bold=True, color=EMPH)
        text(s, x + 0.8, y + 3.3, cw - 1.6, 3.0, [[(t, {"bold": True, "size": 18})], [(sub, {"size": 15, "color": INK2})]],
             space=0.15)

    rx, rw = 35.6, R - 35.6
    picture(s, MEDIA / "web_map_wide.jpg", rx, 5.8, rw, 14.6)
    caption(s, rx, 20.6, rw, "웹앱 지도 — 서울 도심 공개 거리 사진이 지나간 464개 링크를 위험·주의·안전으로 색칠")
    picture(s, MEDIA / "web_map_zoom.jpg", rx, 22.0, rw, 12.6)
    caption(s, rx, 34.8, rw, "확대 — 종로3가~세운상가 골목. 큰길은 안전, 그 사이 이면도로가 위험 · huggingface.co/spaces/YunSangNam/roadpush")


def page2(s, n):
    set_title(s, "2. 아이디어명 / 3. 기술요약", 2)
    header(s, L, 5.6, "2. 아이디어명")
    text(s, L + 0.95, 7.1, CW - 1, 1.4, [[("RoadPush — 보행공간 단절 위험도", {"bold": True}),
                                          ("   걷는 사람이 차도로 밀려나는(push) 구간을 찾는다 · 2026 AI 라이프 아이디어 챌린지 수상 후보작",
                                           {"color": INK2})]], size=20)

    header(s, L, 9.4, "문제 정의")
    probs = [("보행자가 차도로 밀려난다", "보도가 끊기거나(보도 소멸) 주정차·적치물·입간판에 막히면 보행자는 차와 같은 길을 걷는다. "
                                       "휠체어·유모차는 지나갈 수조차 없다."),
             ("점검은 민원·순찰에 의존", "지자체 보도 점검은 사람이 걸으며 눈으로 보고, '얼마나 걷기 어려운가'를 구간끼리 비교할 정량 지표가 없다."),
             ("사진으로 m를 재기는 부정확", "원근 때문에 픽셀 폭을 실제 폭으로 바꾸기 어렵다 → 원래 보도 폭 대비 '비율'로 잰다."),
             ("핵심 질문", "\"어느 도로 구간부터 단속·정비해야 하는가?\"")]
    y = 11.0
    for i, (t, d) in enumerate(probs):
        last = i == len(probs) - 1
        badge(s, L, y + 0.05, "?" if last else str(i + 1), d=1.1, size=15, fill=RED if last else INK2)
        text(s, L + 1.6, y, 26.5, 3.4, [[(t, {"bold": True, "size": 19})], [(d, {"size": 17, "color": INK2})]], space=0.12)
        y += 3.45 if not last else 0
    px, pw, gap = 33.2, (R - 33.2 - 0.7) / 2, 0.7
    picture(s, MLY / "frames" / "p_frame_001837.jpg", px, 9.6, pw, 13.6)
    picture(s, MLY / "frames" / "p_frame_001556.jpg", px + pw + gap, 9.6, pw, 13.6)
    caption(s, px, 23.4, pw, "보도 없이 적치물·사람·차가 섞이는 골목")
    caption(s, px + pw + gap, 23.4, pw, "보도 끝까지 차가 붙어 좁아진 교차로")

    header(s, L, 24.6, "3. 기술요약")
    facts = [f"20종 분할 · 장애물 mIoU {n['obs_miou']}", f"비율 오차 {n['ratio_mae']} · 보도 소멸 F1 {n['vanish_f1']}",
             f"{n['links']}개 링크 · 서버 0대"]
    items = [("분할", "거리 사진을 경량 분할 모델(RF-DETR-Seg, Apache 2.0)로 보도면 9종 · 장애물 11종으로 픽셀 단위 구분한다."),
             ("측정", "가로줄마다 '지나갈 수 있는 폭 ÷ 원래 보도 폭'(유효보행폭 비율)의 병목을 재고, "
                     "보도 소멸·폭 협소·공사·단차·보차 혼용을 판정한다."),
             ("우선순위", "GPS로 보행 링크에 붙여 0~100점 위험도를 내고, AI 점검 메모와 함께 "
                       "서버 없이 도는 지도 웹앱으로 담당자에게 제공한다.")]
    cw, gx, y0, h = (CW - 2 * 0.8) / 3, 0.8, 26.3, 9.9
    for i, (t, d) in enumerate(items):
        x = L + i * (cw + gx)
        box(s, x, y0, cw, h, fill=EMPH_BG)
        badge(s, x + 0.8, y0 + 0.8, str(i + 1), d=1.4, size=19)
        text(s, x + 2.6, y0 + 0.8, cw - 3.2, 1.4, [t], size=22, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.8, y0 + 2.8, cw - 1.6, 4.6, [d], size=19, space=0.2)
        box(s, x + 0.8, y0 + h - 2.3, cw - 1.6, 1.5, fill=WHITE)
        text(s, x + 0.8, y0 + h - 2.3, cw - 1.6, 1.5, [facts[i]], size=17, bold=True, color=EMPH,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def page3(s, n):
    set_title(s, "4. 최종목표 / 5. 개발 내용 및 결과 ① AI 접근방법", 3)
    header(s, L, 5.6, "4. 최종목표")
    text(s, L + 0.95, 7.1, CW - 1, 2.8, [
        "ㅇ 걸으며 눈으로 하던 보도 점검을, 거리 사진만으로 도로 링크 단위 정량 지표(유효보행폭 비율 · 단절 유형 · 위험도)로 바꾼다.",
        "ㅇ 서울 도심 실제 거리 사진으로 '차도로 밀려나는 구간'과 조치 우선순위를 자동 산출하고, 누구나 쓰는 웹앱으로 공개한다."],
         size=19, space=0.25)

    header(s, L, 10.3, "5. 개발 내용 및 결과 — AI 접근방법", w=60)
    steps = [("거리 사진 + GPS", "폰 촬영 · 공개 거리사진\n360° 사진은 진행 방향 변환"),
             ("RF-DETR-Seg 분할", "보도면 9종 · 장애물 11종\n픽셀 단위 마스크"),
             ("유효보행폭 비율", "가로줄별 통과 폭 ÷ 보도 폭\n가장 좁은 곳(병목)"),
             ("링크 위험도", "0~100점 · 단절 유형\n보차 혼용 · 휠체어 기준"),
             ("담당자 웹앱", "지도 · 조치 우선순위\nAI 점검 메모")]
    bw, gap, by, bh = 10.6, (CW - 5 * 10.6) / 4, 12.0, 4.3
    for i, (t, d) in enumerate(steps):
        x = L + i * (bw + gap)
        core = i in (1, 2)
        box(s, x, by, bw, bh, fill=EMPH if core else EMPH_BG, line=None if core else EMPH)
        text(s, x + 0.4, by + 0.45, bw - 0.8, 1.4, [t], size=20, bold=True, align=PP_ALIGN.CENTER,
             color=WHITE if core else INK)
        text(s, x + 0.4, by + 1.95, bw - 0.8, 2.1, d.split("\n"), size=15, align=PP_ALIGN.CENTER,
             color=WHITE if core else INK2, space=0.05)
        if i < 4:
            arrow(s, x + bw + 0.15, by + bh / 2, x + bw + gap - 0.15, by + bh / 2)
    # 학습(오프라인) · 운영
    ty, th = 17.2, 2.7
    tw = 2 * bw + gap
    box(s, L, ty, tw, th, fill=WHITE, line=INK2, dash=True)
    text(s, L + 0.5, ty + 0.25, tw - 1, th - 0.5, [
        [("학습 (오프라인, RTX 3090)", {"bold": True, "size": 16})],
        [("AI Hub 인도보행 영상 보도면 9.6천 · 장애물 1.5만 장 + SAM 3 teacher가 빠진 라벨 14만 개 보강", {"size": 15, "color": INK2})]],
         space=0.1)
    arrow(s, L + bw + gap / 2 + bw / 2, ty - 0.05, L + bw + gap / 2 + bw / 2, by + bh + 0.1, dash=True)
    ox = L + 2 * (bw + gap)
    ow = R - ox
    box(s, ox, ty, ow, th, fill=WHITE, line=INK2, dash=True)
    text(s, ox + 0.5, ty + 0.25, ow - 1, th - 0.5, [
        [("운영 (서버 없음)", {"bold": True, "size": 16})],
        [("ONNX로 변환해 사용자 브라우저에서 추론 · Qwen3-VL-8B가 측정값과 대표 사진을 근거로 원인·권고 메모 작성", {"size": 15, "color": INK2})]],
         space=0.1)

    # 분할 예시 + 비율 개념
    picture(s, FIG / "fig_masks.png", L, 20.7, 38.6)
    px = L + 38.6 + 0.9
    pw = R - px
    box(s, px, 20.7, pw, 15.7, fill=CARD)
    text(s, px + 0.7, 21.2, pw - 1.4, 1.3, ["유효보행폭 비율 (가로줄마다 측정)"], size=19, bold=True)
    # 개념도: 보도 띠 · 장애물 · 통과 폭
    sx, sy, sw = px + 1.2, 23.4, pw - 2.4
    box(s, sx, sy, sw, 2.2, fill=rgb("BFE3C6"), shape=MSO_SHAPE.RECTANGLE)
    box(s, sx + sw * 0.62, sy - 0.2, sw * 0.24, 2.6, fill=RED, shape=MSO_SHAPE.RECTANGLE)
    text(s, sx + sw * 0.62, sy + 0.4, sw * 0.24, 1.4, ["장애물"], size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    text(s, sx + 0.3, sy + 0.4, sw * 0.55, 1.4, ["통과 가능 w"], size=14, bold=True, color=rgb("14532D"))
    arrow(s, sx, sy + 3.0, sx + sw, sy + 3.0, color=INK2, width=1.5)
    arrow(s, sx + sw, sy + 3.0, sx, sy + 3.0, color=INK2, width=1.5)
    text(s, sx, sy + 3.2, sw, 1.1, ["원래 보도 폭 W"], size=14, color=INK2, align=PP_ALIGN.CENTER)
    text(s, px + 0.7, 28.4, pw - 1.4, 7.6, [
        [("비율 = w ÷ W", {"bold": True, "size": 21, "color": EMPH}), ("  가장 좁은 가로줄(병목)", {"size": 17, "color": INK2})],
        [("· 보행자는 장애물로 세지 않음 (조치 대상 아님)", {"size": 16})],
        [("· 보도 소멸 · 보행 공간 위 차량 · 비율 < 0.3 → 밀려남", {"size": 16})],
        [("· 휠체어 기준: 통과 비율 0.7 · 단차 가중", {"size": 16})],
        [("· 링크 점수 = 100 × [0.4(1−하위10% 비율) + 0.3·기준 미달 비중 + 0.3·유형 가중]", {"size": 16})]],
         space=0.35)


def page4(s, n):
    set_title(s, "5. 개발 내용 및 결과 ② 결과 해석", 4)
    header(s, L, 5.6, "모델 성능", w=31, sub="AI Hub 검증셋 · 촬영 폴더 단위 분리 · 같은 150장 비교")
    cats = ["SAM 3\nzero-shot", "Grounding DINO\n+ SAM 2", "RoadPush\n파인튜닝"]
    bar_chart(s, L, 7.3, 14.9, 11.0, "유효보행폭 비율 오차 (낮을수록 좋음)", cats, [0.109, 0.200, 0.038], "0.000", False)
    bar_chart(s, L + 15.6, 7.3, 14.9, 11.0, "장애물 mIoU (높을수록 좋음)", cats, [0.518, 0.373, 0.701], "0.000", True)

    rows = [("전체 검증셋 지표", "값"),
            (f"유효보행폭 비율 오차 (보도면 {n['n_surf']}장)", f"{n['ratio_mae']}  (±0.1 이내 {n['within']})"),
            ("보도 소멸 판정 F1", n["vanish_f1"]),
            (f"장애물 mIoU ({n['n_obs']}장, 희소 클래스 보강 v2)", n["obs_miou"]),
            ("보도 IoU / 보도면 9종 mIoU", f"{n['sidewalk_iou']} / {n['surf_miou']}")]
    tbl = s.shapes.add_table(len(rows), 2, Cm(L), Cm(18.8), Cm(30.5), Cm(5.6)).table
    tbl.columns[0].width, tbl.columns[1].width = Cm(19.5), Cm(11.0)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid(); cell.fill.fore_color.rgb = EMPH_BG if i == 0 else WHITE
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.text = v
            r = cell.text_frame.paragraphs[0].runs[0]
            r.font.size, r.font.color.rgb = Pt(16), INK
            r.font.name = FONT_B if i == 0 or j == 1 else FONT

    header(s, L, 25.2, "결과 해석 — 모델")
    text(s, L + 0.95, 26.8, 29.5, 9.6, [
        [("범용 AI만으로는 부족: ", {"bold": True}),
         (f"zero-shot은 점자블록·가로수 보호틀·입간판을 놓쳐 비율 오차가 {n['x_better']}~5배. 한국 보도 데이터 파인튜닝이 핵심.", {})],
        [("가벼워서 현장 적용 가능: ", {"bold": True}),
         ("서버·GPU 없이 브라우저에서 사진 한 장을 수 초 안에 분석 — 지자체 PC에 설치 없이 바로 쓴다.", {})],
        [("LLM은 교차검증 후 사용: ", {"bold": True}),
         ("멀티모달 LLM 단독 판정은 차량 오경보 29% → 분할 결과와 일치한 사실만 점검 메모에 쓴다.", {})]],
         size=17, space=0.35)

    rx = 35.6
    rw = R - rx
    header(s, rx, 5.6, "서울 도심 실사진 적용", w=rw, sub=f"{n['frames']}장 · {n['links']}개 링크 · {n['km']} km")
    picture(s, FIG / "fig_push_map.png", rx, 7.3, 17.4, 15.3)
    sx, sw = rx + 18.2, R - rx - 18.2
    g = n["grades"]
    chips = [("위험", g["위험"], RED), ("주의", g["주의"], AMBER), ("안전", g["안전"], GREEN)]
    text(s, sx, 7.3, sw, 1.2, ["링크 등급"], size=17, bold=True)
    for i, (k, v, c) in enumerate(chips):
        y = 8.8 + i * 1.75
        box(s, sx, y + 0.2, 0.9, 0.9, fill=c, shape=MSO_SHAPE.OVAL)
        text(s, sx + 1.4, y, sw - 1.4, 1.3, [[(f"{k} ", {"size": 18}), (f"{v}개", {"size": 20, "bold": True})]],
             anchor=MSO_ANCHOR.MIDDLE)
    box(s, sx, 14.4, sw, 7.0, fill=rgb("FCEBEB"))
    text(s, sx + 0.6, 14.9, sw - 1.2, 6.2, [
        [("차도로 밀려나는 링크", {"size": 16, "bold": True})],
        [(f"{n['push_links']}개", {"size": 36, "bold": True, "color": RED})],
        [(f"{n['push_km']} km · 측정 링크의 {n['push_links'] / n['links']:.0%}", {"size": 16, "color": INK2})],
        [(f"보차 혼용 {n['shared']}개 링크", {"size": 16, "color": INK2})]], space=0.15)

    header(s, rx, 23.4, "결과 해석 — 현장", w=rw)
    text(s, rx + 0.95, 25.0, rw - 1, 11.4, [
        [("밀려남은 큰길이 아니라 골목에서: ", {"bold": True}),
         ("밀려나는 링크 164개 중 108개(66%)가 이면도로·골목. 이면도로·골목 링크의 절반 이상이 밀려나는 반면 큰길은 22%.", {})],
        [("주된 원인은 보도 소멸: ", {"bold": True}),
         ("밀려나는 사진의 91%가 '보도가 끊겨 차도·골목으로 합류'. 그다음이 보행 공간을 차지한 차량(9%).", {})],
        [("담당자가 바로 쓸 수 있는 형태: ", {"bold": True}),
         ("링크별 점수·유형·대표 사진·AI 메모가 붙어, 순찰 대상과 보도 신설·보행자 우선도로 후보를 함께 고를 수 있다.", {})]],
         size=17, space=0.35)


def page5(s, n):
    set_title(s, "6. 개발 과정에서 애로점 / 7. 기대효과", 5)
    lw = 30.5
    header(s, L, 5.6, "6. 개발 과정에서 애로점 → 해결", w=lw)
    issues = [("라벨 결측", "보도면 데이터엔 장애물이, 장애물 데이터엔 보도면이 없음",
               "두 전문 모델 결합 + SAM 3 의사 라벨 통합 모델 비교 (정확도 −3점 대신 비용 절반)"),
              ("평가 누수", "연속 프레임이 학습·검증에 섞여 점수가 부풀려짐", "촬영 폴더 단위로 학습/검증 분리"),
              ("지표 희석", "노점·입간판이 있어도 비율이 1.0으로 나옴", "평균 대신 가장 좁은 곳(병목) · 마스크 틈 메움"),
              ("원근", "픽셀 폭을 실제 m로 바꾸기 어려움", "비율을 기본 지표로, 1.2 m는 카메라 높이 가정 보조 지표"),
              ("데이터 접근", "지도앱 로드뷰는 약관상 사용 불가", "Mapillary(CC BY-SA) 보행 사진 + 360° 사진 진행 방향 변환"),
              ("배포", "무료 GPU 서버 호스팅 불가", "ONNX 변환 → 브라우저 안에서 추론하는 정적 웹앱")]
    y, h, gy = 7.3, 4.6, 0.25
    for t, prob, sol in issues:
        box(s, L, y, lw, h, fill=CARD)
        text(s, L + 0.6, y + 0.35, 6.2, h - 0.7, [t], size=18, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        text(s, L + 6.9, y + 0.35, lw - 7.5, h - 0.7, [[(prob, {"size": 16, "color": INK2})],
                                                    [("→ " + sol, {"size": 16, "bold": True})]],
             anchor=MSO_ANCHOR.MIDDLE, space=0.15)
        y += h + gy

    rx = 35.6
    rw = R - rx
    header(s, rx, 5.6, "한계 및 개선방안", w=rw)
    lims = [("사진이 있는 길만 측정, 촬영 시점 혼재(2015~2025)", "청소차·순찰차·배달 이륜차 카메라로 정기 촬영, 시점별 변화 비교"),
            ("실제 폭(m)은 보조 지표 수준", "깊이 추정 · 카메라 높이 보정으로 m 단위 병행"),
            ("계단·공사 등 희소 클래스, 야간·우천 미검증", "현장 데이터 보강 · 지자체와 실측 검증"),
            ("점수 가중치는 연구자 설정값", "담당자 피드백으로 가중치 보정, 조치 전·후 재측정")]
    y, h = 7.3, 3.3
    for lim, fix in lims:
        box(s, rx, y, rw, h, fill=CARD)
        text(s, rx + 0.6, y + 0.3, rw - 1.2, h - 0.6, [[("한계  ", {"size": 16, "bold": True, "color": RED}), (lim, {"size": 17})],
                                                   [("개선  ", {"size": 16, "bold": True, "color": GREEN}), (fix, {"size": 17, "color": INK2})]],
             anchor=MSO_ANCHOR.MIDDLE, space=0.15)
        y += h + 0.25

    header(s, rx, 22.0, "7. 기대효과", w=rw)
    effs = [("지자체", "민원·순찰 의존 점검 → 촬영 후 자동 리포트. 링크별 단속·정비 우선순위와 예산 배분 근거"),
            ("보행약자", "휠체어 기준(폭 0.7·단차) 토글로 지나갈 수 없는 구간을 정량화 — 디지털 포용 기초자료"),
            ("확장 · 개방", "경량 모델이라 청소차·순찰차 카메라로 넓히기 쉬움. 코드 Apache-2.0 공개, 웹앱은 설치 없이 접속")]
    y, h = 23.7, 4.05
    for t, d in effs:
        box(s, rx, y, rw, h, fill=EMPH_BG)
        text(s, rx + 0.6, y + 0.3, 6.0, h - 0.6, [t], size=19, bold=True, color=EMPH, anchor=MSO_ANCHOR.MIDDLE)
        text(s, rx + 6.8, y + 0.3, rw - 7.4, h - 0.6, [d], size=17, anchor=MSO_ANCHOR.MIDDLE)
        y += h + 0.25


def build(out: Path) -> Path:
    n = numbers()
    prs = Presentation(str(ROOT / "docs/templates/성과보고서_양식_PPT_5p.pptx"))
    for fn, s in zip((page1, page2, page3, page4, page5), prs.slides):
        fn(s, n)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out


def main(argv: list[str] | None = None) -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", type=Path, default=ROOT / "dist" / "RoadPush_성과보고서.pptx")
    ap.add_argument("--pdf", action="store_true")
    args = ap.parse_args(argv)
    out = build(args.out)
    print("→", out)
    if args.pdf:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(out.parent), str(out)],
                       check=False, capture_output=True)
        print("→", out.with_suffix(".pdf"))


if __name__ == "__main__":
    main()
