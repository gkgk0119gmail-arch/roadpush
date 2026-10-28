"""RoadPush 성과보고서 (KEIT PPT 양식 기반, 5쪽) → docs/report/RoadPush_성과보고서.pptx

먼저 scripts/media/crop_scenes.py. 폰트는 양식과 같은 Pretendard(없으면 PowerPoint가 대체 폰트로 표시 — 제출은 PDF 병행).
"""
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
INFO = ROOT / "outputs/media/img"          # crop_scenes.py 결과
MEDIA = ROOT / "docs/media"
FIG = ROOT / "docs/figures"
TMP = ROOT / "outputs/media/crops"; TMP.mkdir(parents=True, exist_ok=True)

INK, INK2, MUTE = "14161B", "4F535B", "8A8E96"
BLUE, BLUE_D, BLUE_BG = "4D9CFE", "2A6FD6", "EEF5FF"
RED, RED_BG = "D6403A", "FCEDEB"
AMB, GRN = "E89A12", "1A9A55"
CARD, LINE, GREY = "F4F6F8", "E2E5EA", "B5B9C0"
F_R, F_S, F_B = "Pretendard", "Pretendard SemiBold", "Pretendard ExtraBold"

prs = Presentation(ROOT / "docs/templates/성과보고서_양식_PPT_5p.pptx")   # KEIT 양식 4p + 결과 해석 1쪽
slides = list(prs.slides)


def rgb(h): return RGBColor.from_string(h)


def set_font(run, size, color=INK, font=F_R, bold=False):
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, font
    f.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {}); rpr.append(el)
        el.set("typeface", font)


def text(sl, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=0, line=1.15, name=None):
    """paras: [[(text, size, color, font), ...], ...]"""
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name: tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        if space and i: p.space_before = Pt(space)
        for t, size, color, font in runs:
            r = p.add_run(); r.text = t
            set_font(r, size, color, font)
    return tb


def T(t, size=24, color=INK, font=F_R): return (t, size, color, font)


def box(sl, x, y, w, h, fill=CARD, line=None, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name=None):
    s = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name: s.name = name
    s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    if line: s.line.color.rgb = rgb(line); s.line.width = Pt(1.25)
    else: s.line.fill.background()
    s.shadow.inherit = False
    st = s._element.find(qn('p:style'))
    if st is not None: s._element.remove(st)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def circle_num(sl, x, y, d, label, fill=BLUE_D, size=24):
    c = box(sl, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    tf = c.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label; set_font(r, size, "FFFFFF", F_B)
    return c


def pic(sl, path, x, y, w, h, pos=(0.5, 0.5), name=None):
    """박스 비율에 맞게 잘라서 넣기 (cover)."""
    im = Image.open(path).convert("RGB")
    W, H = im.size; ta = w / h
    if W / H > ta:
        nw = int(H * ta); x0 = int((W - nw) * pos[0]); im = im.crop((x0, 0, x0 + nw, H))
    else:
        nh = int(W / ta); y0 = int((H - nh) * pos[1]); im = im.crop((0, y0, W, y0 + nh))
    out = TMP / f"{Path(path).stem}_{w:.2f}x{h:.2f}_{pos[0]}_{pos[1]}.jpg"
    im.save(out, quality=90)
    p = sl.shapes.add_picture(str(out), Inches(x), Inches(y), Inches(w), Inches(h))
    if name: p.name = name
    return p


def arrow(sl, x1, y1, x2, y2, color=GREY, wt=3):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color); c.line.width = Pt(wt)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}); ln.append(tail)
    return c


def head(sl, x, y, w, label, sub=None):
    runs = [T(label, 28, INK, F_S)]
    if sub: runs.append(T("   " + sub, 20, MUTE, F_R))
    return text(sl, x, y, w, 0.6, [runs])


def clean(sl, title=None, page=None):
    for sh in list(sl.shapes):
        if sh.name in ("TextBox 5", "TextBox 6"):
            sh._element.getparent().remove(sh._element)
        elif sh.name == "TextBox 2":
            sh.width = Inches(19)
            if title:
                p = sh.text_frame.paragraphs[0]; rs = p.runs
                rs[0].text = title
                for r in rs[1:]: r.text = ""
        elif sh.name == "TextBox 12" and page:
            sh.text_frame.paragraphs[0].runs[0].text = str(page)


L, R_EDGE = 1.36, 25.3          # 본문 좌우 경계
CW = R_EDGE - L                 # 23.94

# ───────────────────────── 1. AI 기술명 ─────────────────────────
s = slides[0]; clean(s)
text(s, L, 2.3, 11.8, 1.4, [[T("RoadPush", 88, INK, F_B)]])
text(s, L, 3.85, 11.8, 0.7, [[T("보행 영상 기반 보행공간 단절 위험도 자동 산출 시스템", 32, INK, F_S)]])
text(s, L, 4.75, 11.6, 1.3, [[T("거리 사진만으로 ‘걷는 사람이 차도로 밀려나는(push) 구간’을 찾아, 도로 구간마다 위험도와 "
                               "조치 우선순위를 보여 주는 지자체 보도 담당자용 AI", 23, INK2)]], line=1.3)
box(s, L, 6.35, 11.6, 2.75, fill=RED_BG, name="밀려남 카드")
text(s, L + 0.5, 6.55, 5.6, 2.3, [[T("8.7", 112, RED, F_B), T(" km", 44, RED, F_B)]], anchor=MSO_ANCHOR.MIDDLE)
text(s, L + 6.1, 6.7, 5.2, 2.1, [[T("걷는 사람이 차도로 밀려나는 길", 26, INK, F_S)],
                                 [T("서울 도심 측정 28.3 km 중 31%", 21, INK2)],
                                 [T("464개 도로 구간 중 164개", 21, INK2)]], anchor=MSO_ANCHOR.MIDDLE, space=6)
kpis = [("0.017", "유효보행폭 비율 오차", "AI Hub 검증 954장"), ("1/3", "범용 AI(SAM 3) 대비 오차", "0.109 → 0.038"),
        ("0.755", "장애물 인식 mIoU", "검증 1,566장 · 11종"), ("서버 0대", "브라우저 안에서 AI 실행", "설치 없이 웹 접속")]
kw, kh = 5.55, 1.75
for i, (v, l1, l2) in enumerate(kpis):
    x = L + (i % 2) * (kw + 0.5); y = 9.45 + (i // 2) * (kh + 0.3)
    box(s, x, y, kw, kh, fill=CARD)
    text(s, x + 0.35, y + 0.2, kw - 0.5, 0.85, [[T(v, 40, BLUE_D, F_B)]])
    text(s, x + 0.35, y + 1.02, kw - 0.5, 0.6, [[T(l1, 19, INK, F_S), T("  " + l2, 16, MUTE)]])
# 탐구참여인력
text(s, L, 13.6, 2.6, 0.5, [[T("탐구참여인력", 20, INK, F_S)]])
tbl = s.shapes.add_table(1, 3, Inches(L + 2.6), Inches(13.52), Inches(9.0), Inches(0.62)).table
for j, (t, w) in enumerate([("남윤상", 1.6), ("연세대학교 건설환경공학과", 3.6), ("기획·데이터·모델·웹앱 (1인)", 3.8)]):
    tbl.columns[j].width = Inches(w)
    c = tbl.cell(0, j); c.fill.solid(); c.fill.fore_color.rgb = rgb(CARD)
    c.margin_left = Inches(0.15); c.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = c.text_frame.paragraphs[0]; r = p.add_run(); r.text = t; set_font(r, 18, INK, F_S if j == 0 else F_R)
# 오른쪽: 실제 웹앱
RX, RW = 13.75, R_EDGE - 13.75
pic(s, MEDIA / "web_map_wide.jpg", RX, 2.35, RW, RW / 1.93, name="웹앱 지도")
text(s, RX, 2.35 + RW / 1.93 + 0.12, RW, 0.45, [[T("웹앱 지도 — 서울 도심 464개 구간을 위험·주의·안전으로 표시", 18, INK2),
                                                T("   huggingface.co/spaces/YunSangNam/roadpush", 18, BLUE_D, F_S)]])
iw = (RW - 0.4) / 2; ih = 3.9
y2 = 9.45
pic(s, MEDIA / "web_detail.jpg", RX, y2, iw, ih, pos=(0.5, 0.0), name="링크 상세")
pic(s, MEDIA / "web_upload.jpg", RX + iw + 0.4, y2, iw, ih, pos=(0.5, 0.3), name="사진 분석")
text(s, RX, y2 + ih + 0.12, iw, 0.45, [[T("링크 상세 — 점수·단절 유형·AI 점검 메모", 18, INK2)]])
text(s, RX + iw + 0.4, y2 + ih + 0.12, iw, 0.45, [[T("사진 분석 — 올린 사진을 브라우저에서 즉시 분할", 18, INK2)]])

# ───────────────────────── 2. 아이디어명 / 기술요약 ─────────────────────────
s = slides[1]; clean(s)
box(s, L, 2.3, CW, 1.45, fill=BLUE_BG)
text(s, L + 0.4, 2.42, 3.2, 0.6, [[T("2. 아이디어명", 24, BLUE_D, F_S)]])
text(s, L + 3.6, 2.42, CW - 4.0, 0.6, [[T("RoadPush — 보행공간 단절 위험도", 28, INK, F_B),
                                       T("   2026 AI 라이프 아이디어 챌린지 수상 후보작", 20, INK2)]])
text(s, L + 3.6, 3.08, CW - 4.0, 0.5, [[T("공공 도로·보행 영상을 AI로 분석해 보행공간 단절 위험도를 산출·시각화하고 사전 예방에 기여 (제안 취지)", 19, INK2)]])
head(s, L, 4.1, 14, "문제 정의", "보도가 있어도 걸을 수 없는 길 — 사람은 차도로 밀려난다")
cw4 = (CW - 3 * 0.5) / 4; ph = 3.55; cy = 4.8
cards = [(INFO / "b_BL.jpg", "보도가 끊긴다", "보도가 사라져 사람·수레·차가\n한 길을 나눠 씀", (0.5, 0.6)),
         (INFO / "f_1.jpg", "차가 보도를 차지한다", "보도블록 위로 차가 드나들어\n사람이 비켜서는 길", (0.5, 0.6)),
         (INFO / "b_TR.jpg", "장애물이 길을 막는다", "입간판·적치물로 지나갈 폭이\n원래의 30% 미만", (0.5, 0.6))]
for i, (img, t1, t2, pos) in enumerate(cards):
    x = L + i * (cw4 + 0.5)
    box(s, x, cy, cw4, 5.3, fill=CARD)
    pic(s, img, x, cy, cw4, ph, pos=pos, name=t1)
    circle_num(s, x + 0.25, cy + 0.25, 0.65, str(i + 1), fill=RED, size=24)
    text(s, x + 0.3, cy + ph + 0.2, cw4 - 0.5, 0.55, [[T(t1, 25, INK, F_S)]])
    text(s, x + 0.3, cy + ph + 0.78, cw4 - 0.5, 0.9, [[T(t2, 18, INK2)]], line=1.2)
x = L + 3 * (cw4 + 0.5)
box(s, x, cy, cw4, 5.3, fill=INK, name="현재 점검 방식")
text(s, x + 0.45, cy + 0.45, cw4 - 0.9, 4.5, [
    [T("지금의 점검", 20, "9FC4FF", F_S)],
    [T("민원·순찰에 의존", 30, "FFFFFF", F_B)],
    [T("사람이 걸으며 눈으로 확인 — 구간끼리 비교할 정량 지표가 없다", 19, "D5D8DE")],
    [T("핵심 질문", 20, "9FC4FF", F_S)],
    [T("“어느 도로 구간부터 단속·정비해야 하는가?”", 24, "FFFFFF", F_S)]], space=10, line=1.25)
head(s, L, 10.45, 14, "3. 기술요약")
sw = (CW - 2 * 0.5) / 3; sy = 11.15; sh_ = 3.3
steps = [("분할", "거리 사진을 경량 분할 모델(RF-DETR-Seg, Apache 2.0)로 보도면 9종·장애물 11종으로 픽셀 단위 구분",
          "장애물 mIoU 0.755"),
         ("측정", "가로줄마다 ‘지나갈 수 있는 폭 ÷ 원래 보도 폭’(유효보행폭 비율)의 병목을 재고 보도 소멸·폭 협소·보차 혼용 등을 판정",
          "비율 오차 0.017 · 소멸 F1 0.915"),
         ("우선순위", "GPS로 도로 구간(링크)에 모아 0~100점 위험도와 AI 점검 메모를 만들고, 서버 없는 지도 웹앱으로 담당자에게 제공",
          "464개 구간 · 브라우저 실행")]
for i, (t1, t2, kp) in enumerate(steps):
    x = L + i * (sw + 0.5)
    box(s, x, sy, sw, sh_, fill="FFFFFF", line=LINE)
    circle_num(s, x + 0.35, sy + 0.35, 0.7, str(i + 1))
    text(s, x + 1.3, sy + 0.38, sw - 1.6, 0.65, [[T(t1, 27, INK, F_S)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.35, sy + 1.25, sw - 0.7, 1.4, [[T(t2, 19, INK2)]], line=1.28)
    text(s, x + 0.35, sy + sh_ - 0.6, sw - 0.7, 0.45, [[T(kp, 19, BLUE_D, F_S)]])
    if i < 2: arrow(s, x + sw + 0.07, sy + sh_ / 2, x + sw + 0.43, sy + sh_ / 2, color=BLUE)

# ───────────────────────── 3. 최종목표 / AI 접근방법 ─────────────────────────
s = slides[2]; clean(s, title="4. 최종목표 / 5. 개발 내용 및 결과 ① AI 접근방법")
box(s, L, 2.3, CW, 1.5, fill=BLUE_BG)
text(s, L + 0.4, 2.45, 3.0, 0.55, [[T("4. 최종목표", 24, BLUE_D, F_S)]])
text(s, L + 3.4, 2.45, CW - 3.8, 1.5, [
    [T("① ", 21, BLUE_D, F_S), T("걸으며 눈으로 하던 보도 점검을, 거리 사진만으로 도로 구간 단위 정량 지표(유효보행폭 비율·단절 유형·위험도)로 바꾼다", 21, INK)],
    [T("② ", 21, BLUE_D, F_S), T("서울 도심 실제 사진으로 ‘차도로 밀려나는 구간’과 조치 우선순위를 자동 산출하고, 누구나 쓰는 웹앱으로 공개한다", 21, INK)]],
    space=6, line=1.25)
head(s, L, 4.4, 20, "5. 개발 내용 및 결과 — AI 접근방법")
flow = [("거리 사진 + GPS", "폰 촬영·공개 거리 사진\n360° 사진은 진행 방향 변환"),
        ("RF-DETR-Seg 분할", "보도면 9종 · 장애물 11종\n픽셀 단위 마스크"),
        ("유효보행폭 비율", "지나갈 폭 ÷ 원래 보도 폭\n가장 좁은 곳(병목)"),
        ("구간 위험도", "0~100점 · 단절 유형\n보차 혼용 · 휠체어 기준"),
        ("담당자 웹앱", "지도 · 조치 우선순위\nAI 점검 메모")]
fw = (CW - 4 * 0.55) / 5; fy = 5.1; fh = 1.95
for i, (t1, t2) in enumerate(flow):
    x = L + i * (fw + 0.55)
    hi = i in (1, 2)
    box(s, x, fy, fw, fh, fill=BLUE_D if hi else "FFFFFF", line=None if hi else LINE)
    text(s, x + 0.3, fy + 0.25, fw - 0.6, 0.55, [[T(t1, 22, "FFFFFF" if hi else INK, F_S)]])
    text(s, x + 0.3, fy + 0.88, fw - 0.6, 1.0, [[T(t2, 16, "DCE8FF" if hi else INK2)]], line=1.2)
    if i < 4: arrow(s, x + fw + 0.07, fy + fh / 2, x + fw + 0.48, fy + fh / 2, color=BLUE)
# 원본 vs AI 분할
mw = 14.6; mh = mw / (2417 / 988)
pic(s, FIG / "fig_masks.png", L, 7.5, mw, mh, name="원본과 AI 분할")
text(s, L, 7.5 + mh + 0.1, mw, 0.45, [[T("원본(위)과 AI 분할(아래) — 초록 보도 · 노랑 점자블록 · 갈색 골목 · 빨강 장애물 · 연보라 보행자", 17, INK2)]])
# 비율 설명 카드
bx, bw = L + mw + 0.55, CW - mw - 0.55
box(s, bx, 7.5, bw, 6.45, fill=CARD, name="비율 정의")
text(s, bx + 0.4, 7.7, bw - 0.8, 0.55, [[T("유효보행폭 비율", 24, INK, F_S), T("  가로줄마다 측정", 17, MUTE)]])
# 도식: 보도 띠 + 장애물
dx, dy, dw = bx + 0.4, 8.5, bw - 0.8
box(s, dx, dy, dw, 0.75, fill="BFE3C9", radius=0.15)
box(s, dx + dw * 0.62, dy, dw * 0.38, 0.75, fill=RED, radius=0.15)
text(s, dx + dw * 0.62, dy, dw * 0.38, 0.75, [[T("장애물", 17, "FFFFFF", F_S)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, dx + 0.15, dy, dw * 0.5, 0.75, [[T("지나갈 수 있는 폭  w", 17, "1D6B35", F_S)]], anchor=MSO_ANCHOR.MIDDLE)
arrow(s, dx, dy + 1.0, dx + dw, dy + 1.0, color=INK2, wt=2)
text(s, dx, dy + 1.08, dw, 0.4, [[T("원래 보도 폭  W", 16, INK2)]], align=PP_ALIGN.CENTER)
text(s, bx + 0.4, 10.15, bw - 0.8, 0.6, [[T("비율 = w ÷ W", 26, BLUE_D, F_B), T("   가장 좁은 줄(병목)이 사진의 값", 16, INK2)]])
rules = ["보행자는 장애물에서 제외 (조치 대상 아님)",
         "바닥이 안 보이는 사진은 ‘측정 불가’로 제외",
         "밀려남: 보도 소멸 · 보행 공간 위 차량 · 비율 < 0.3",
         "휠체어 기준: 통과 비율 0.7 · 단차 가중",
         "원근 때문에 m 대신 비율 (1.2 m는 보조 지표)"]
text(s, bx + 0.4, 10.9, bw - 0.8, 2.9, [[T("· " + r, 17, INK)] for r in rules], space=5, line=1.15)
text(s, L, 14.05, CW, 0.45, [[T("학습", 16, BLUE_D, F_S), T("  RTX 3090 · AI Hub 인도보행 영상(보도면 9.6천 장 · 장애물 2.4만 장) · SAM 3·Grounding DINO는 비교·라벨링 teacher로만", 16, INK2),
                             T("     운영", 16, BLUE_D, F_S), T("  ONNX로 브라우저 추론 · Qwen3-VL-8B 점검 메모", 16, INK2)]])

# ───────────────────────── 4. 결과 해석 ─────────────────────────
s = slides[3]; clean(s, title="5. 개발 내용 및 결과 ② 결과 해석", page=4)
LW = 12.4
head(s, L, 2.3, LW, "모델 성능", "AI Hub 검증셋 · 촬영 폴더 단위로 학습과 분리")


def bar(sl, x, y, w, h, title, vals, fmt, maxv):
    cd = CategoryChartData()
    cd.categories = ["GDINO+SAM 2\n(범용)", "SAM 3\n(범용)", "RoadPush\n(파인튜닝)"]
    cd.add_series(title, vals)
    gf = sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = True
    ch.chart_title.text_frame.text = title
    for r in ch.chart_title.text_frame.paragraphs[0].runs: set_font(r, 18, INK, F_S)
    pl = ch.plots[0]; pl.gap_width = 70; pl.has_data_labels = True
    dl = pl.data_labels; dl.number_format = fmt; dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(18); dl.font.bold = True; dl.font.name = F_S
    ser = pl.series[0]
    for i, c in enumerate([GREY, GREY, BLUE_D]):
        pt = ser.points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(c)
    va = ch.value_axis; va.visible = False; va.has_major_gridlines = False
    va.maximum_scale = maxv; va.minimum_scale = 0
    ca = ch.category_axis; ca.tick_labels.font.size = Pt(15); ca.tick_labels.font.name = F_R
    ca.tick_labels.font.color.rgb = rgb(INK2); ca.format.line.color.rgb = rgb(LINE)
    ca.has_major_gridlines = False
    return gf


bar(s, L, 2.95, 6.0, 5.1, "지나갈 폭 비율 오차 (낮을수록 좋음)", (0.200, 0.109, 0.038), "0.000", 0.24)
bar(s, L + 6.4, 2.95, 6.0, 5.1, "장애물 인식 mIoU (높을수록 좋음)", (0.373, 0.518, 0.701), "0.000", 0.85)
text(s, L, 8.1, LW, 0.4, [[T("같은 검증 사진 150장 · 같은 클래스 기준 비교", 16, MUTE)]])
rows = [("지표 (검증셋 전체)", "값", "의미"),
        ("유효보행폭 비율 오차", "0.017", "954장 · 97%가 ±0.1 이내"),
        ("보도 소멸 판정 F1", "0.915", "보도가 끊긴 사진 판별"),
        ("보도면 mIoU", "0.670", "보도·차도·골목 등 9종"),
        ("장애물 mIoU", "0.755", "1,566장 · 희소 클래스 보강")]
tb = s.shapes.add_table(len(rows), 3, Inches(L), Inches(8.7), Inches(LW), Inches(2.9)).table
for j, w in enumerate([4.6, 1.8, 6.0]): tb.columns[j].width = Inches(w)
for i, row in enumerate(rows):
    for j, t in enumerate(row):
        c = tb.cell(i, j); c.fill.solid(); c.fill.fore_color.rgb = rgb(BLUE_BG if i == 0 else ("FFFFFF" if i % 2 else CARD))
        c.margin_left = Inches(0.18); c.vertical_anchor = MSO_ANCHOR.MIDDLE
        r = c.text_frame.paragraphs[0].add_run(); r.text = t
        set_font(r, 18, BLUE_D if (j == 1 and i) else INK, F_S if (i == 0 or j == 1) else F_R)
text(s, L, 11.95, LW, 0.55, [[T("결과 해석 — 모델", 22, INK, F_S)]])
text(s, L, 12.55, LW, 2.0, [
    [T("범용 AI만으로는 부족 ", 18, INK, F_S), T("zero-shot은 차도·골목, 볼라드·입간판을 거의 구분하지 못함 → 한국 보도 데이터 파인튜닝이 핵심", 18, INK2)],
    [T("작은 모델로 현장 적용 ", 18, INK, F_S), T("3,600만 파라미터 모델이 8억 파라미터 SAM 3보다 오차 1/3 — 서버 없이 브라우저에서 실행", 18, INK2)],
    [T("LLM은 근거가 있을 때만 ", 18, INK, F_S), T("단독 판정 시 차량 오경보 29% → 분할 결과와 교차검증한 사실만 메모에 사용", 18, INK2)]],
    space=6, line=1.2)
# 오른쪽: 서울 도심
RX = L + LW + 0.7; RW = R_EDGE - RX
head(s, RX, 2.3, RW, "서울 도심 실사진 적용", "7,559장 · 464개 구간 · 28.3 km")
mw, mh = 6.9, 6.9 / 1.279
box(s, RX, 2.95, mw, mh, fill="FCFCFB", line=LINE)
pic(s, INFO / "map.png", RX + 0.05, 3.0, mw - 0.1, mh - 0.1, name="밀려나는 구간 지도")
sx = RX + mw + 0.4; sw2 = R_EDGE - sx
box(s, sx, 2.95, sw2, 2.35, fill=RED_BG)
text(s, sx + 0.3, 3.1, sw2 - 0.5, 2.1, [[T("차도로 밀려나는 구간", 17, INK, F_S)], [T("164개", 44, RED, F_B)], [T("8.7 km · 측정의 31%", 17, INK2)]], line=1.1)
for i, (lab, n, col) in enumerate([("위험", "101", RED), ("주의", "69", AMB), ("안전", "294", GRN)]):
    y = 5.55 + i * 0.75
    circle_num(s, sx, y + 0.08, 0.42, "", fill=col)
    text(s, sx + 0.6, y, sw2 - 0.6, 0.6, [[T(lab, 19, INK2), T(f"  {n}개", 21, INK, F_S)]], anchor=MSO_ANCHOR.MIDDLE)
text(s, sx, 7.85, sw2, 0.6, [[T("보차 혼용", 19, INK2), T("  149개", 21, INK, F_S)]], anchor=MSO_ANCHOR.MIDDLE)
gy = 2.95 + mh + 0.35; gw = (RW - 2 * 0.3) / 3; gh = 2.45
for i, (img, cap) in enumerate([(INFO / "b_TL.jpg", "창경궁로5가길 · 92점"), (INFO / "a_TR.jpg", "을지로33길 · 98점"), (INFO / "b_BR.jpg", "동호로38길 · 95점")]):
    x = RX + i * (gw + 0.3)
    pic(s, img, x, gy, gw, gh, pos=(0.5, 0.65), name=cap)
    text(s, x, gy + gh + 0.08, gw, 0.4, [[T(cap, 16, INK2)]])
text(s, RX, gy + gh + 0.65, RW, 0.55, [[T("결과 해석 — 현장", 22, INK, F_S)]])
text(s, RX, gy + gh + 1.25, RW, 2.0, [
    [T("큰길은 안전, 그 사이 골목이 끊긴다 ", 18, INK, F_S), T("— 종로3가~세운상가 일대처럼 구간 단위로 골목별 조치 순서를 정할 수 있음", 18, INK2)],
    [T("보도블록이 깔린 길도 위험할 수 있다 ", 18, INK, F_S), T("— 차가 드나드는 보차 혼용 149개 구간을 별도 표시", 18, INK2)]],
    space=6, line=1.2)

# ───────────────────────── 5. 애로점 / 기대효과 / 한계 ─────────────────────────
s = slides[4]; clean(s, page=5)
LW = 12.6
head(s, L, 2.3, LW, "6. 개발 과정에서 애로점 → 해결")
issues = [("라벨이 서로 빠져 있음", "AI Hub 보도면 데이터엔 장애물, 장애물 데이터엔 보도면 라벨이 없음",
           "전문 모델 2개 결합(채택) vs SAM 3 의사 라벨 14만 개 통합 모델 — 정확도 −3점, 비용 절반"),
          ("검증 점수 부풀림", "연속 프레임이 학습·검증에 섞여 성능이 과대평가", "촬영 폴더 단위로 학습/검증 분리"),
          ("장애물이 있어도 비율 1.0", "노점·입간판이 평균에 희석, 마스크 틈으로 분리", "평균 대신 가장 좁은 곳(병목) + 통과 여유"),
          ("실사진 오판", "가게 진열창·근접 사진이 ‘보도 소멸’로 집계", "바닥이 안 보이면 ‘측정 불가’로 제외"),
          ("LLM 환각", "8B 모델도 없는 차량을 있다고 답함(오경보 29%)", "닫힌 질문 + 분할 결과와 교차검증한 사실만 사용"),
          ("데이터·배포", "지도앱 로드뷰는 약관상 불가, 무료 서버 호스팅 중단", "Mapillary(CC BY-SA)+OpenStreetMap, 브라우저 추론 웹앱")]
iy, ih_ = 3.0, 1.82
for i, (t1, t2, t3) in enumerate(issues):
    y = iy + i * (ih_ + 0.1)
    box(s, L, y, LW, ih_, fill=CARD)
    text(s, L + 0.35, y + 0.18, LW - 0.7, 0.5, [[T(t1, 21, RED, F_S), T("   " + t2, 17, INK2)]])
    text(s, L + 0.35, y + 0.82, LW - 0.7, 0.9, [[T("→ ", 19, BLUE_D, F_B), T(t3, 19, INK, F_S)]], line=1.2)
RX = L + LW + 0.6; RW = R_EDGE - RX
head(s, RX, 2.3, RW, "7. 기대효과")
eff = [("지자체", "민원·순찰 의존 점검 → 촬영 후 자동 리포트. 구간별 단속·정비 우선순위와 예산 배분 근거"),
       ("보행약자", "휠체어 기준(통과 폭 0.7·단차 가중)으로 지나갈 수 없는 구간을 정량화 — 디지털 포용 정책 기초자료"),
       ("확장", "청소차·순찰차·배달 이륜차 카메라에 그대로 적용(경량 모델, 서버 불필요). 정기 재분석으로 상습 구간 탐지")]
ey, eh = 3.0, 1.95
for i, (t1, t2) in enumerate(eff):
    y = ey + i * (eh + 0.2)
    box(s, RX, y, RW, eh, fill=BLUE_BG)
    text(s, RX + 0.35, y + 0.2, 2.3, eh - 0.4, [[T(t1, 24, BLUE_D, F_B)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, RX + 2.6, y + 0.2, RW - 2.95, eh - 0.4, [[T(t2, 18, INK)]], anchor=MSO_ANCHOR.MIDDLE, line=1.25)
head(s, RX, 9.65, RW, "한계 및 개선방안")
lim = [("한계", "개선방안"),
       ("사진이 지나간 길만 측정, 촬영 시점 혼재(2015~2025)", "공공 차량 카메라로 정기 촬영 → 시점별 변화·조치 전후 비교"),
       ("정량 검증은 AI Hub 검증셋 기준", "지자체와 현장 실측·담당자 판정으로 실사진 검증"),
       ("점수 가중치는 연구자 설정값", "담당자 피드백으로 가중치 보정"),
       ("계단·공사 등 희소 유형, 야간·우천 미검증", "현장 데이터 보강 후 재학습")]
tb = s.shapes.add_table(len(lim), 2, Inches(RX), Inches(10.3), Inches(RW), Inches(4.1)).table
tb.columns[0].width = Inches(RW * 0.5); tb.columns[1].width = Inches(RW * 0.5)
for i, row in enumerate(lim):
    for j, t in enumerate(row):
        c = tb.cell(i, j); c.fill.solid(); c.fill.fore_color.rgb = rgb((RED_BG if j == 0 else BLUE_BG) if i == 0 else ("FFFFFF" if i % 2 else CARD))
        c.margin_left = Inches(0.18); c.margin_right = Inches(0.12); c.vertical_anchor = MSO_ANCHOR.MIDDLE
        r = c.text_frame.paragraphs[0].add_run(); r.text = t
        set_font(r, 18 if i == 0 else 16, (RED if j == 0 else BLUE_D) if i == 0 else INK, F_S if i == 0 else F_R)

out = ROOT / "docs/report/RoadPush_성과보고서.pptx"; out.parent.mkdir(exist_ok=True)
prs.save(out)
print("saved", out)
